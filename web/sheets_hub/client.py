from __future__ import annotations

import io
import threading
import time
import xml.etree.ElementTree as ET
from dataclasses import replace
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlparse

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from sheets_hub.auth import AuthError, account_label, load_nextcloud_credentials
from sheets_hub.calendar_sheet import (
    is_calendar_matrix,
    is_lock_text,
    lock_is_fresh,
    make_lock_text,
    parse_sheet_rows,
)
from sheets_hub.config import (
    KIND_INFO,
    KIND_RECORDS,
    SheetRef,
    companion_info_titles,
    is_info_ref,
    is_info_title,
    parse_spreadsheet_id,
    prefer_sheet_title,
    requested_sheet_titles,
)
from sheets_hub.models import Record
from sheets_hub.registry import (
    DEFAULT_REGISTRY_SHEET,
    refs_from_registry_rows,
    refs_to_registry_rows,
    row_looks_like_registry_header,
)
from sheets_hub.split import address_key, service_key
from sheets_hub.ssl_setup import configure_tls, session_verify_target

MAX_SHEET_ROWS = 38
MAX_INFO_SHEET_ROWS = 120
MAX_CALENDAR_WITH_NOTES_ROWS = 80
_COLOR_CACHE: dict[str, tuple[float, dict[tuple[int, int], str]]] = {}
_COLOR_CACHE_TTL_SEC = 900
_PREFERRED_CALENDAR_SHEET: dict[str, str] = {}
_WB_LOCKS: dict[str, threading.Lock] = {}
_WB_LOCKS_GUARD = threading.Lock()


class SheetsError(Exception):
    pass


def contrast_fg(bg_hex: str) -> str:
    raw = (bg_hex or "").lstrip("#")
    if len(raw) != 6:
        return "#202124"
    try:
        r, g, b = int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)
    except ValueError:
        return "#202124"
    lum = (0.299 * r + 0.587 * g + 0.114 * b) / 255.0
    return "#202124" if lum > 0.55 else "#ffffff"


def soften_fill(bg_hex: str, *, fallback: str = "#1b5e20") -> str:
    """Приглушить слишком яркие заливки для экрана."""
    raw = (bg_hex or "").lstrip("#")
    if len(raw) != 6:
        return fallback
    try:
        r, g, b = int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)
    except ValueError:
        return fallback
    lum = (0.299 * r + 0.587 * g + 0.114 * b) / 255.0
    if g >= 160 and g >= r + 30 and g >= b + 30 and lum >= 0.42:
        return fallback
    if lum <= 0.50:
        return f"#{r:02x}{g:02x}{b:02x}"
    target = 0.40
    factor = target / lum if lum else 1.0
    r = max(0, min(255, int(round(r * factor))))
    g = max(0, min(255, int(round(g * factor))))
    b = max(0, min(255, int(round(b * factor))))
    return f"#{r:02x}{g:02x}{b:02x}"


def enrich_records_with_sheet_colors(
    records: list[Record],
    color_map: dict[tuple[int, int], str],
) -> None:
    if not records or not color_map:
        return
    by_row: dict[int, list[tuple[int, str]]] = {}
    for (row, col), hex_c in color_map.items():
        by_row.setdefault(row, []).append((col, hex_c))
    for row in by_row:
        by_row[row].sort(key=lambda item: item[0])

    for record in records:
        if record.layout == "calendar":
            col = record.col_index.get("Клиент")
            if not col:
                continue
            bg = color_map.get((record.row, int(col)))
        else:
            col = record.col_index.get("Текст")
            bg = color_map.get((record.row, int(col))) if col else None
            if not bg:
                row_colors = by_row.get(record.row) or []
                if row_colors:
                    bg = row_colors[0][1]
        if bg:
            record.values["_bg"] = bg


def _color_bounds_from_records(records: list[Record]) -> tuple[int, int, list[int]]:
    cols: set[int] = set()
    rows: list[int] = []
    for record in records:
        if record.layout == "calendar":
            col = record.col_index.get("Клиент")
        else:
            col = record.col_index.get("Текст")
        if col:
            try:
                cols.add(int(col))
            except (TypeError, ValueError):
                pass
        if record.row:
            try:
                rows.append(int(record.row))
            except (TypeError, ValueError):
                pass
    if not rows:
        return 1, MAX_CALENDAR_WITH_NOTES_ROWS, sorted(cols)
    row_min = max(1, min(rows) - 1)
    row_max = min(MAX_INFO_SHEET_ROWS, max(rows) + 1)
    return row_min, row_max, sorted(cols)


def _row_limit_for_source(source: SheetRef, sheet_title: str = "") -> int:
    title = (sheet_title or source.sheet or "").strip()
    if is_info_ref(source) or is_info_title(title):
        return MAX_INFO_SHEET_ROWS
    return MAX_CALENDAR_WITH_NOTES_ROWS


def _parts_for_source(source: SheetRef, available: list[str] | None = None) -> list[SheetRef]:
    requested = requested_sheet_titles(source.sheet)
    if requested is not None and len(requested) == 1 and not is_info_title(requested[0]):
        calendar = replace(source, sheet=requested[0])
        parts = [calendar]
        if available:
            for title in companion_info_titles(available, requested[0]):
                parts.append(replace(source, sheet=title, kind=KIND_INFO))
        return parts
    if requested is not None and len(requested) > 1:
        return [replace(source, sheet=title) for title in requested]

    titles = available or []
    if not titles:
        return [source]
    calendar_title = prefer_sheet_title(titles, source.sheet)
    if is_info_title(calendar_title):
        non_info = [title for title in titles if not is_info_title(title)]
        if non_info:
            calendar_title = prefer_sheet_title(non_info, source.sheet)
    parts = [replace(source, sheet=calendar_title)]
    for title in companion_info_titles(titles, calendar_title):
        parts.append(replace(source, sheet=title, kind=KIND_INFO))
    return parts


def _apply_map(raw: dict[str, str], mapping: dict[str, str]) -> dict[str, str]:
    if not mapping:
        return dict(raw)
    out = dict(raw)
    for canonical, alias in mapping.items():
        if alias in raw and (canonical not in out or not str(out.get(canonical) or "").strip()):
            out[canonical] = raw[alias]
    return out


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.replace("\u202f", " ").strip()
    if hasattr(value, "strftime"):
        try:
            return value.strftime("%d.%m.%Y")
        except Exception:
            pass
    return str(value).strip()


def _fill_to_hex(fill: PatternFill | None) -> str | None:
    if fill is None:
        return None
    color = None
    try:
        if fill.fgColor and fill.fgColor.type == "rgb" and fill.fgColor.rgb:
            color = str(fill.fgColor.rgb)
        elif fill.start_color and fill.start_color.type == "rgb" and fill.start_color.rgb:
            color = str(fill.start_color.rgb)
    except Exception:
        return None
    if not color:
        return None
    raw = color[-6:] if len(color) >= 6 else color
    if len(raw) != 6:
        return None
    try:
        r, g, b = int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16)
    except ValueError:
        return None
    if r >= 250 and g >= 250 and b >= 250:
        return None
    if r + g + b < 8:
        return None
    return f"#{r:02x}{g:02x}{b:02x}"


def _workbook_lock(path_key: str) -> threading.Lock:
    with _WB_LOCKS_GUARD:
        lock = _WB_LOCKS.get(path_key)
        if lock is None:
            lock = threading.Lock()
            _WB_LOCKS[path_key] = lock
        return lock


def _pick_registry_title(
    wanted: str,
    available: list[str],
    header_by_title: dict[str, list] | None = None,
) -> str | None:
    want = (wanted or DEFAULT_REGISTRY_SHEET).strip().lower()
    for title in available:
        if title.lower() == want:
            return title
    if header_by_title:
        for title, header in header_by_title.items():
            if row_looks_like_registry_header(header):
                return title
    for title in available:
        if want and want in title.lower():
            return title
    if available:
        return available[0]
    return None


class SheetsClient:
    """Клиент календарей на Nextcloud: .xlsx через WebDAV."""

    def __init__(self, credentials_path: Path, *, interactive: bool = False) -> None:
        del interactive
        configure_tls()
        if not credentials_path.exists():
            raise SheetsError(
                "Нет credentials.json для Nextcloud.\n"
                "Нажмите «Таблицы» → «JSON ключа» и укажите файл с url / username / app_password."
            )
        self._credentials_path = credentials_path
        self.read_only_public = False
        self.read_notes: list[str] = []
        self.auth_kind = "nextcloud"
        self._etag_cache: dict[str, str] = {}
        self._connect()

    def _connect(self) -> None:
        configure_tls()
        try:
            creds = load_nextcloud_credentials(self._credentials_path)
        except AuthError as exc:
            raise SheetsError(str(exc)) from exc
        self._base_url = creds["url"].rstrip("/")
        self._username = creds["username"]
        self._password = creds["app_password"]
        self.service_email = account_label(self._credentials_path, creds)
        self._session = requests.Session()
        self._session.auth = (self._username, self._password)
        self._session.verify = session_verify_target()
        self._session.headers["User-Agent"] = "SheetsHub-Nextcloud/1.0"
        retry = Retry(
            total=2,
            connect=2,
            read=2,
            status=2,
            backoff_factor=0.4,
            status_forcelist=(429, 500, 502, 503, 504),
            allowed_methods=False,
            raise_on_status=False,
        )
        adapter = HTTPAdapter(max_retries=retry, pool_connections=4, pool_maxsize=4)
        self._session.mount("https://", adapter)
        self._session.mount("http://", adapter)
        self.auth_kind = "nextcloud"

    def _dav_root(self) -> str:
        user = quote(self._username, safe="")
        return f"{self._base_url}/remote.php/dav/files/{user}/"

    def _dav_url(self, remote_path: str) -> str:
        path = parse_spreadsheet_id(remote_path)
        path = path.lstrip("/")
        parts = [quote(part, safe="") for part in path.split("/") if part]
        return self._dav_root() + "/".join(parts)

    def _raise_http(self, resp: requests.Response, *, what: str) -> None:
        code = resp.status_code
        if code == 401:
            raise SheetsError(
                "Nextcloud отклонил вход (401). Проверьте username и app_password "
                "(Настройки → Безопасность → пароли приложений)."
            )
        if code == 403:
            raise SheetsError(f"Нет доступа к файлу ({what}). Проверьте права шаринга.")
        if code == 404:
            raise SheetsError(f"Файл не найден на Nextcloud: {what}")
        if code == 429:
            raise SheetsError(
                "Nextcloud временно ограничил запросы (429 Too Many Requests).\n"
                "Подождите 1–2 минуты и не жмите «Обновить» подряд.\n"
                f"Файл: {what}"
            )
        text = (resp.text or "")[:200].strip()
        raise SheetsError(f"Nextcloud HTTP {code} ({what}): {text or resp.reason}")

    def _request_with_retry(
        self,
        method: str,
        url: str,
        *,
        what: str,
        retries: int = 3,
        **kwargs,
    ) -> requests.Response:
        last: requests.Response | None = None
        for attempt in range(max(1, retries)):
            resp = self._session.request(method, url, **kwargs)
            last = resp
            if resp.status_code != 429:
                return resp
            # Nextcloud rate-limit — подождать и повторить.
            time.sleep(1.5 * (attempt + 1))
        assert last is not None
        return last

    def probe_api(self, spreadsheet_id: str = "") -> bool:
        try:
            target = self._dav_url(spreadsheet_id) if spreadsheet_id else self._dav_root()
            resp = self._request_with_retry(
                "PROPFIND",
                target,
                what=spreadsheet_id or "/",
                headers={"Depth": "0"},
                timeout=(4, 12),
                retries=2,
            )
            ok = resp.status_code in (207, 200)
            self._api_ok = ok
            return ok
        except Exception:
            self._api_ok = False
            return False

    def _download_bytes(self, remote_path: str) -> tuple[bytes, str]:
        url = self._dav_url(remote_path)
        resp = self._request_with_retry(
            "GET",
            url,
            what=parse_spreadsheet_id(remote_path),
            timeout=(5, 60),
            retries=3,
        )
        if resp.status_code >= 400:
            self._raise_http(resp, what=parse_spreadsheet_id(remote_path))
        etag = (resp.headers.get("ETag") or resp.headers.get("OC-Etag") or "").strip()
        if etag:
            self._etag_cache[parse_spreadsheet_id(remote_path)] = etag
        return resp.content, etag

    def _upload_bytes(self, remote_path: str, data: bytes, *, etag: str = "") -> None:
        url = self._dav_url(remote_path)
        headers: dict[str, str] = {
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }
        if etag:
            headers["If-Match"] = etag
        resp = self._request_with_retry(
            "PUT",
            url,
            what=parse_spreadsheet_id(remote_path),
            data=data,
            headers=headers,
            timeout=(5, 90),
            retries=3,
        )
        if resp.status_code == 412:
            raise SheetsError(
                "Файл изменился на сервере, пока вы редактировали.\n"
                "Обновите календарь и повторите."
            )
        if resp.status_code >= 400:
            self._raise_http(resp, what=parse_spreadsheet_id(remote_path))
        new_etag = (resp.headers.get("ETag") or resp.headers.get("OC-Etag") or "").strip()
        if new_etag:
            self._etag_cache[parse_spreadsheet_id(remote_path)] = new_etag

    def _ensure_parent_dirs(self, remote_path: str) -> None:
        path = parse_spreadsheet_id(remote_path).lstrip("/")
        parts = path.split("/")[:-1]
        if not parts:
            return
        cur = ""
        for part in parts:
            cur = f"{cur}/{part}" if cur else part
            url = self._dav_url(cur)
            self._session.request("MKCOL", url, timeout=(4, 15))

    def _normalize_xlsx_path(self, remote_path: str) -> str:
        path = parse_spreadsheet_id(remote_path)
        if not path:
            raise SheetsError("Не указан путь к файлу на Nextcloud")
        if not path.lower().endswith((".xlsx", ".xlsm")):
            if "." not in Path(path).name:
                path = path + ".xlsx"
            else:
                raise SheetsError(
                    f"Поддерживаются файлы .xlsx (Nextcloud Office / Excel): {path}"
                )
        return path

    def _load_workbook(self, remote_path: str, *, data_only: bool = False):
        path = self._normalize_xlsx_path(remote_path)
        raw, _etag = self._download_bytes(path)
        try:
            return load_workbook(io.BytesIO(raw), data_only=data_only), path
        except Exception as exc:
            raise SheetsError(f"Не удалось открыть таблицу «{path}»: {exc}") from exc

    def _save_workbook(self, remote_path: str, wb, *, etag: str = "") -> None:
        buf = io.BytesIO()
        wb.save(buf)
        self._ensure_parent_dirs(remote_path)
        self._upload_bytes(remote_path, buf.getvalue(), etag=etag)

    def _match_sheet_title(self, available: list[str], wanted: str) -> str | None:
        raw = (wanted or "").strip()
        if not raw:
            return None
        for title in available:
            if title.lower() == raw.lower():
                return title
        for title in available:
            if raw.lower() in title.lower() or title.lower() in raw.lower():
                return title
        return None

    def _open_worksheet(self, ref: SheetRef):
        path = ref.normalized_id()
        if not path or path.upper().startswith("PASTE_"):
            raise SheetsError(f"В «{ref.name}» не указан путь к файлу Nextcloud")
        with _workbook_lock(path):
            wb, path = self._load_workbook(path)
            available = list(wb.sheetnames)
            if not available:
                raise SheetsError(f"В «{ref.name}» нет листов")
            match = self._match_sheet_title(available, ref.sheet)
            if match is None and len(available) == 1:
                match = available[0]
            if match is None:
                raise SheetsError(
                    f"В «{ref.name}» нет листа «{ref.sheet}». "
                    f"В файле есть: {', '.join(available)}."
                )
            return wb, wb[match], path

    def _sheet_values(self, worksheet, *, max_rows: int | None = None) -> list[list[str]]:
        limit = MAX_CALENDAR_WITH_NOTES_ROWS if max_rows is None else max(1, int(max_rows))
        rows: list[list[str]] = []
        max_col = min(worksheet.max_column or 1, 60)
        for row in worksheet.iter_rows(min_row=1, max_row=limit, max_col=max_col, values_only=True):
            values = [_cell_to_str(cell) for cell in row]
            while values and not values[-1]:
                values.pop()
            rows.append(values)
        while rows and not any(rows[-1]):
            rows.pop()
        return rows

    def get_headers(self, ref: SheetRef) -> list[str]:
        wb, worksheet, _path = self._open_worksheet(ref)
        try:
            rows = self._sheet_values(worksheet, max_rows=_row_limit_for_source(ref))
        finally:
            wb.close()
        if not rows:
            return []
        if is_calendar_matrix(rows):
            return []
        return [cell.strip() for cell in rows[0] if cell.strip()]

    def _records_from_rows(self, source: SheetRef, rows: list[list[str]]) -> list[Record]:
        spreadsheet_id = source.normalized_id()
        if not rows:
            return []
        limit = _row_limit_for_source(source)
        rows = rows[:limit]

        parsed = parse_sheet_rows(rows, source, spreadsheet_id)
        if parsed is not None:
            return parsed

        headers = [cell.strip() for cell in rows[0]]
        col_index = {header: idx + 1 for idx, header in enumerate(headers) if header}

        records: list[Record] = []
        for offset, row in enumerate(rows[1:], start=2):
            if not any(cell.strip() for cell in row):
                continue
            raw = {
                headers[i]: row[i] if i < len(row) else ""
                for i in range(len(headers))
                if headers[i]
            }
            values = _apply_map(raw, source.map)
            if source.service:
                key = service_key(values) or "Тип услуги"
                if not str(values.get(key, "")).strip():
                    values[key] = source.service
            if source.address:
                key = address_key(values) or "Адрес"
                if not str(values.get(key, "")).strip():
                    values[key] = source.address
            records.append(
                Record(
                    source_name=source.name,
                    spreadsheet_id=spreadsheet_id,
                    sheet=source.sheet,
                    row=offset,
                    values=values,
                    sheet_headers=headers,
                    col_index=col_index,
                    map=source.map,
                    origin_values=dict(values),
                    kind=KIND_INFO if is_info_ref(source) else KIND_RECORDS,
                )
            )
        return records

    def fetch_source(self, source: SheetRef) -> list[Record]:
        wb, worksheet, path = self._open_worksheet(source)
        try:
            rows = self._sheet_values(worksheet, max_rows=_row_limit_for_source(source))
            part = replace(source, sheet=worksheet.title, spreadsheet_id=path)
            return self._records_from_rows(part, rows)
        finally:
            wb.close()

    def expand_source(self, source: SheetRef) -> list[SheetRef]:
        available = self.list_sheets(source.spreadsheet_id)
        parts = _parts_for_source(source, available)
        if not parts:
            raise SheetsError(f"В «{source.name}» нет листов")
        return parts

    def _extract_sheet_colors(
        self,
        worksheet,
        *,
        row_start: int,
        row_end: int,
        cols: list[int] | None,
    ) -> dict[tuple[int, int], str]:
        colors: dict[tuple[int, int], str] = {}
        wanted = set(cols or [])
        max_col = min(worksheet.max_column or 1, 60)
        for row in range(max(1, row_start), max(row_start, row_end) + 1):
            for col in range(1, max_col + 1):
                if wanted and col not in wanted:
                    continue
                cell = worksheet.cell(row=row, column=col)
                hex_c = _fill_to_hex(cell.fill)
                if hex_c:
                    colors[(row, col)] = hex_c
        return colors

    def _attach_sheet_colors(self, records: list[Record], *, force: bool = False) -> None:
        if not records:
            return
        by_sheet: dict[tuple[str, str], list[Record]] = {}
        for record in records:
            key = (record.spreadsheet_id, record.sheet or "")
            by_sheet.setdefault(key, []).append(record)

        for (sid, title), group in by_sheet.items():
            if not sid or not title:
                continue
            cache_key = f"{sid}|{title}"
            now = time.time()
            cached = _COLOR_CACHE.get(cache_key)
            if not force and cached and (now - cached[0]) < _COLOR_CACHE_TTL_SEC:
                enrich_records_with_sheet_colors(group, cached[1])
                continue
            try:
                row_start, row_end, cols = _color_bounds_from_records(group)
                with _workbook_lock(sid):
                    wb, path = self._load_workbook(sid)
                    try:
                        match = self._match_sheet_title(list(wb.sheetnames), title)
                        if not match:
                            continue
                        colors = self._extract_sheet_colors(
                            wb[match],
                            row_start=row_start,
                            row_end=row_end,
                            cols=cols or None,
                        )
                    finally:
                        wb.close()
                if colors:
                    merged = dict(cached[1]) if cached else {}
                    merged.update(colors)
                    _COLOR_CACHE[cache_key] = (now, merged)
                    enrich_records_with_sheet_colors(group, merged)
                elif cached:
                    enrich_records_with_sheet_colors(group, cached[1])
            except Exception:
                if cached:
                    enrich_records_with_sheet_colors(group, cached[1])

    def apply_cached_colors(self, records: list[Record]) -> None:
        if not records:
            return
        for record in records:
            sid, title = record.spreadsheet_id, record.sheet or ""
            if not sid or not title:
                continue
            cached = _COLOR_CACHE.get(f"{sid}|{title}")
            if not cached:
                continue
            enrich_records_with_sheet_colors([record], cached[1])

    def fetch_all(
        self,
        sources: list[SheetRef],
        *,
        include_colors: bool = True,
        fast: bool = False,
    ) -> tuple[list[Record], list[str]]:
        del fast
        records: list[Record] = []
        errors: list[str] = []
        self.read_only_public = False
        self.read_notes = []
        for source in sources:
            if source.is_placeholder():
                continue
            try:
                batch: list[Record] = []
                for part in self.expand_source(source):
                    batch.extend(self.fetch_source(part))
                if include_colors:
                    self._attach_sheet_colors(batch)
                for item in batch:
                    if item.layout == "calendar" and item.sheet:
                        _PREFERRED_CALENDAR_SHEET[item.spreadsheet_id] = item.sheet
                        break
                records.extend(batch)
            except Exception as exc:
                msg = str(exc) if isinstance(exc, SheetsError) else str(exc)
                errors.append(f"{source.name}: {msg}")
        return records, errors

    def update_cell(self, record: Record, field: str, value: str, *, confirm: bool = True) -> None:
        try:
            header = record.header_for(field)
        except KeyError as exc:
            raise SheetsError(f"В листе нет колонки «{field}»") from exc
        col = record.col_index[header]
        cell_value = value
        if record.layout == "calendar" and field == "Клиент":
            cell_value = value.strip() or "запись"
        path = self._normalize_xlsx_path(record.spreadsheet_id)
        with _workbook_lock(path):
            raw, etag = self._download_bytes(path)
            wb = load_workbook(io.BytesIO(raw))
            try:
                match = self._match_sheet_title(list(wb.sheetnames), record.sheet)
                if not match:
                    raise SheetsError(f"Лист «{record.sheet}» не найден в {path}")
                ws = wb[match]
                ws.cell(row=record.row, column=int(col), value=cell_value)
                self._save_workbook(path, wb, etag=etag)
            finally:
                wb.close()
        if confirm:
            got = self.read_cell(record, field)
            expect = cell_value.strip()
            if got != expect and not (
                record.layout == "calendar" and field == "Клиент" and got in (expect, "запись")
            ):
                raise SheetsError(
                    f"Запись не подтвердилась. Ожидали «{expect}», в файле «{got}»."
                )

    def read_cell(self, record: Record, field: str) -> str:
        try:
            header = record.header_for(field)
        except KeyError as exc:
            raise SheetsError(f"В листе нет колонки «{field}»") from exc
        col = record.col_index[header]
        path = self._normalize_xlsx_path(record.spreadsheet_id)
        with _workbook_lock(path):
            wb, _path = self._load_workbook(path, data_only=True)
            try:
                match = self._match_sheet_title(list(wb.sheetnames), record.sheet)
                if not match:
                    raise SheetsError(f"Лист «{record.sheet}» не найден")
                value = wb[match].cell(row=record.row, column=int(col)).value
                return _cell_to_str(value)
            finally:
                wb.close()

    def acquire_calendar_lock(
        self,
        record: Record,
        *,
        previous_hint: str | None = None,
    ) -> tuple[str, str]:
        if previous_hint is not None and not (
            is_lock_text(previous_hint) and lock_is_fresh(previous_hint)
        ):
            current = previous_hint
        else:
            current = self.read_cell(record, "Клиент")
            if is_lock_text(current) and lock_is_fresh(current):
                raise SheetsError(
                    "Этот слот сейчас заполняет другой оператор.\n"
                    "Подождите немного или нажмите «Обновить»."
                )
        lock_text, _token = make_lock_text()
        self.update_cell(record, "Клиент", lock_text, confirm=False)
        return current, lock_text

    def assert_calendar_lock(self, record: Record, lock_text: str) -> None:
        current = self.read_cell(record, "Клиент")
        if current == lock_text:
            return
        if is_lock_text(current) and lock_is_fresh(current):
            raise SheetsError(
                "Слот перехватил другой оператор. Ваша запись не сохранена.\n"
                "Обновите календарь."
            )
        raise SheetsError(
            "Ячейка изменилась, пока вы её редактировали.\n"
            f"Сейчас в таблице: «{current or '(пусто)'}».\n"
            "Сохранение отменено — обновите календарь."
        )

    def append_row(self, dest: SheetRef, values: dict[str, str]) -> None:
        path = self._normalize_xlsx_path(dest.spreadsheet_id)
        with _workbook_lock(path):
            raw, etag = self._download_bytes(path)
            wb = load_workbook(io.BytesIO(raw))
            try:
                match = self._match_sheet_title(list(wb.sheetnames), dest.sheet)
                if not match and len(wb.sheetnames) == 1:
                    match = wb.sheetnames[0]
                if not match:
                    raise SheetsError(f"Лист «{dest.sheet}» не найден")
                ws = wb[match]
                headers = [
                    _cell_to_str(c)
                    for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
                ]
                if not any(headers):
                    raise SheetsError(
                        f"В «{dest.name}» нет заголовков в первой строке. "
                        "Сначала подпишите колонки в таблице Nextcloud."
                    )
                row = [values.get(header, "") for header in headers]
                ws.append(row)
                self._save_workbook(path, wb, etag=etag)
            finally:
                wb.close()

    def delete_row(self, record: Record) -> None:
        path = self._normalize_xlsx_path(record.spreadsheet_id)
        with _workbook_lock(path):
            raw, etag = self._download_bytes(path)
            wb = load_workbook(io.BytesIO(raw))
            try:
                match = self._match_sheet_title(list(wb.sheetnames), record.sheet)
                if not match:
                    raise SheetsError(f"Лист «{record.sheet}» не найден")
                wb[match].delete_rows(record.row)
                self._save_workbook(path, wb, etag=etag)
            finally:
                wb.close()

    def list_sheets(self, url_or_id: str) -> list[str]:
        path = self._normalize_xlsx_path(url_or_id)
        with _workbook_lock(path):
            wb, _path = self._load_workbook(path)
            try:
                return list(wb.sheetnames)
            finally:
                wb.close()

    def list_calendar_sheet_titles(self, spreadsheet_id: str) -> list[str]:
        try:
            titles = self.list_sheets(spreadsheet_id)
        except Exception:
            return []
        return [title for title in titles if title and not is_info_title(title)]

    def preferred_calendar_sheet(self, spreadsheet_id: str) -> str:
        try:
            sid = parse_spreadsheet_id(spreadsheet_id)
        except ValueError:
            return ""
        return _PREFERRED_CALENDAR_SHEET.get(sid, "")

    def set_preferred_calendar_sheet(self, spreadsheet_id: str, sheet: str) -> None:
        try:
            sid = parse_spreadsheet_id(spreadsheet_id)
        except ValueError:
            return
        title = (sheet or "").strip()
        if title:
            _PREFERRED_CALENDAR_SHEET[sid] = title
        else:
            _PREFERRED_CALENDAR_SHEET.pop(sid, None)

    def pull_table_registry(self, spreadsheet_id: str, sheet_title: str = "") -> list[SheetRef]:
        sid = parse_spreadsheet_id(spreadsheet_id)
        title = (sheet_title or DEFAULT_REGISTRY_SHEET).strip() or DEFAULT_REGISTRY_SHEET
        with _workbook_lock(sid):
            wb, path = self._load_workbook(sid)
            try:
                available = list(wb.sheetnames)
                header_by_title: dict[str, list] = {}
                for name in available[:8]:
                    try:
                        first = [
                            _cell_to_str(c)
                            for c in next(wb[name].iter_rows(min_row=1, max_row=1, values_only=True))
                        ]
                        if first:
                            header_by_title[name] = first
                    except Exception:
                        continue
                actual = _pick_registry_title(title, available, header_by_title)
                if actual is None:
                    return []
                rows = self._sheet_values(wb[actual], max_rows=500)
                return refs_from_registry_rows(rows)
            finally:
                wb.close()

    def push_table_registry(
        self,
        spreadsheet_id: str,
        tables: list[SheetRef],
        sheet_title: str = "",
    ) -> str:
        title = (sheet_title or DEFAULT_REGISTRY_SHEET).strip() or DEFAULT_REGISTRY_SHEET
        sid = parse_spreadsheet_id(spreadsheet_id)
        rows = refs_to_registry_rows(tables)
        with _workbook_lock(sid):
            path = self._normalize_xlsx_path(sid)
            try:
                raw, etag = self._download_bytes(path)
                wb = load_workbook(io.BytesIO(raw))
                created = False
            except SheetsError as exc:
                # Новый файл — только если его правда нет; 429/401 не считаем «пустым».
                if "не найден" not in str(exc).lower() and "404" not in str(exc):
                    raise
                wb = Workbook()
                etag = ""
                created = True
                sid = path
            try:
                available = list(wb.sheetnames)
                actual = self._match_sheet_title(available, title)
                if actual is None:
                    if created and available:
                        ws = wb.active
                        ws.title = title
                        actual = title
                    else:
                        ws = wb.create_sheet(title)
                        actual = ws.title
                else:
                    ws = wb[actual]
                if ws.max_row and ws.max_row > 0:
                    ws.delete_rows(1, ws.max_row)
                for row in rows:
                    ws.append(list(row))
                self._save_workbook(sid, wb, etag=etag)
                return actual
            finally:
                wb.close()

    def list_remote_xlsx(self, folder: str = "") -> list[str]:
        """Список .xlsx в папке Nextcloud."""
        path = parse_spreadsheet_id(folder).strip("/") if folder else ""
        url = self._dav_url(path) if path else self._dav_root()
        resp = self._session.request(
            "PROPFIND",
            url,
            headers={"Depth": "1"},
            timeout=(5, 30),
        )
        if resp.status_code >= 400:
            self._raise_http(resp, what=path or "/")
        out: list[str] = []
        try:
            root = ET.fromstring(resp.content)
        except ET.ParseError:
            return out
        ns = {"d": "DAV:"}
        for href_el in root.findall(".//d:href", ns):
            href = unquote((href_el.text or "").strip())
            if not href.lower().endswith((".xlsx", ".xlsm")):
                continue
            idx = href.find("/remote.php/dav/files/")
            if idx >= 0:
                rest = href[idx:].split("/", 5)
                rel = rest[5] if len(rest) >= 6 else href
            else:
                rel = href
            rel = rel.lstrip("/")
            if rel:
                out.append(rel)
        return sorted(set(out))
